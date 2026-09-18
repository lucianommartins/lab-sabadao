# -*- coding: utf-8 -*-
# Copyright 2026 Google LLC
# Evaluation Suite: gdpval
# Description: GDPval (OpenAI Real Economic Knowledge-Work Deliverables Benchmark)

"""gbench native built-in runner for gdpval (Economic & Knowledge Work).

GDPval: 220 economically-valuable knowledge-work tasks (the public "gold" subset). 125 of
them (57%) take an attached spreadsheet/PDF/document as input, and every task ships an
EXPERT REFERENCE DELIVERABLE (`deliverable_files`) plus a long grading rubric.

**Metric: pairwise win-rate, the canonical GDPval measure.** GDPval is scored by comparing a
model's deliverable against the human-expert reference in a blind pairwise judgement (win /
tie / loss), and reporting the win-rate. gbench reproduces that: the model's answer and the
expert reference are shown to the judge as A/B (order arbitrary), judged against the rubric,
position-swapped to cancel order bias, and the headline is the win-rate (wins + 0.5*ties over
the scored tasks). The judge is gbench's standard Gemini cascade (a gbench convention).

**This is not OpenAI's GDPval number and says so on every result** (`leaderboard_comparable`
False): a text `/v1` endpoint cannot emit an `.xlsx`/`.docx`/`.pdf`, so the expert reference
FILE is rendered to text and the model competes on rendered content only - it can never match
the file's live formulas/formatting. A task whose reference cannot be rendered to text, or
where the judge is unavailable, is excluded from the win-rate (not counted as a loss); an
empty/failed model deliverable IS counted, as a loss.

Attachments are fetched with `hf_hub_download` and converted per type; `load_dataset` returns
only the file PATHS. See docs/evals/gdpval.md.

Sampling: temperature defaults to 0.0 (greedy) for no-think runs and 1.0 (the model's
shipped `generation_config.json`) with `--thinking`; the measurement is in
base.DEFAULT_TEMPERATURE. Override it for a whole run with
`--temperature`, or for this suite alone with `GBENCH_GDPVAL_TEMPERATURE`,
which takes precedence over both. LLM-judge grading is pinned at 0.0 and is not
affected by either.
"""

import json
import logging
import os
import re
from typing import Any, Dict, List, Optional, Tuple

from .base import (DEFAULT_JUDGE_MODEL, gemini_required_skip, judge_config,
                   judge_generate_cascade, run_eval_suite, strip_thinking_tags)
from .sampling import stratified_sample

logger = logging.getLogger(__name__)

PILLAR = "Economic & Knowledge Work"
DOCS_URL = "docs/evals/gdpval.md"
_REPO = "openai/gdpval"

#: Attachment budget per file. GDPval inputs run to hundreds of KB of spreadsheet; without a
#: cap a single task can blow the context window on one CSV dump.
_MAX_CHARS_PER_FILE = int(os.getenv("GBENCH_GDPVAL_MAX_FILE_CHARS", "20000"))
_MAX_PDF_PAGES = int(os.getenv("GBENCH_GDPVAL_MAX_PDF_PAGES", "10"))

#: Per-deliverable clip handed to the pairwise judge (model answer and rendered reference each),
#: and the rubric clip. Keeps a single pairwise prompt inside a sane context budget.
_MAX_JUDGE_CHARS = int(os.getenv("GBENCH_GDPVAL_MAX_JUDGE_CHARS", "24000"))
_MAX_RUBRIC_CHARS = int(os.getenv("GBENCH_GDPVAL_MAX_RUBRIC_CHARS", "8000"))


# --------------------------------------------------------------------------- #
# attachments
# --------------------------------------------------------------------------- #
def _fetch(rel_path: str) -> Optional[str]:
    """Resolve one dataset-relative attachment path to a local file (cached)."""
    try:
        from huggingface_hub import hf_hub_download
        return hf_hub_download(repo_id=_REPO, filename=rel_path, repo_type="dataset")
    except Exception as e:                                          # noqa: BLE001
        logger.warning("gdpval: could not fetch %s (%s)", rel_path, e)
        return None


def render_attachment(rel_path: str) -> Tuple[Optional[str], Optional[str]]:
    """Convert one attachment to text. Returns (text, note).

    Exactly one of the two is set. `note` carries the reason a file could not be rendered,
    which is attached to the prompt verbatim - a task whose input silently went missing
    would look like a model failure, which is the thing this suite exists to avoid.
    Images are handled separately by :func:`attachment_parts`.
    """
    name = os.path.basename(rel_path)
    local = _fetch(rel_path)
    if local is None:
        return None, f"[attachment '{name}' could not be downloaded]"
    return render_local_file(local, name)


def render_local_file(local: str, name: str) -> Tuple[Optional[str], Optional[str]]:
    """Convert an already-materialised file to text. Returns (text, note).

    Split out of `render_attachment` so callers that already hold the BYTES can reuse the
    same converters. A text-gradeable plugin variant needs this: its attachments are embedded in the
    export rather than fetched by path, and until 2026-08-19 they were dropped entirely -
    the model was asked to work from .xlsx/.docx/.pdf documents it never received, and
    unsurprisingly replied that it could not produce the deliverable.
    """
    ext = name.rsplit(".", 1)[-1].lower() if "." in name else ""
    try:
        if ext in ("xlsx", "xlsm"):
            import pandas as pd
            sheets = pd.read_excel(local, sheet_name=None)
            out = []
            for sname, df in sheets.items():
                out.append(f"### sheet: {sname}  ({df.shape[0]} rows x {df.shape[1]} cols)\n"
                           + df.to_csv(index=False))
            return _clip("\n".join(out), name), None
        if ext == "pdf":
            import fitz
            doc = fitz.open(local)
            pages = [doc[i].get_text() for i in range(min(doc.page_count, _MAX_PDF_PAGES))]
            extra = ("" if doc.page_count <= _MAX_PDF_PAGES
                     else f"\n[... {doc.page_count - _MAX_PDF_PAGES} further pages omitted]")
            return _clip("\n\n".join(pages) + extra, name), None
        if ext == "docx":
            try:
                import docx
                d = docx.Document(local)
                parts = [p.text for p in d.paragraphs if p.text.strip()]
                for t in d.tables:
                    parts.append("\n".join(" | ".join(c.text for c in row.cells)
                                           for row in t.rows))
                return _clip("\n".join(parts), name), None
            except Exception:
                # python-docx raises XMLSyntaxError on some real-but-nonstandard docx
                # (measured 2026-08-20: 2 gdpval files). The raw word/document.xml is valid
                # enough to read - unzip it and strip tags so the model still gets the
                # content instead of "[could not be read]".
                import zipfile
                xml = zipfile.ZipFile(local).read("word/document.xml").decode(
                    "utf-8", "replace")
                text = re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", xml)).strip()
                if not text:
                    raise
                return _clip(text, name), None
        if ext == "pptx":
            from pptx import Presentation
            pr = Presentation(local)
            slides = []
            for i, s in enumerate(pr.slides, 1):
                txt = "\n".join(sh.text for sh in s.shapes if hasattr(sh, "text") and sh.text)
                slides.append(f"### slide {i}\n{txt}")
            return _clip("\n".join(slides), name), None
        if ext in ("txt", "csv", "md", "json"):
            with open(local, encoding="utf-8", errors="replace") as fh:
                return _clip(fh.read(), name), None
        if ext == "zip":
            import zipfile
            names = zipfile.ZipFile(local).namelist()
            return _clip("archive contents:\n" + "\n".join(names), name), None
        if ext in ("wav", "mp3", "m4a", "flac"):
            import soundfile as sf
            i = sf.info(local)
            return None, (f"[attachment '{name}' is {i.duration:.0f}s of audio; this suite "
                          f"does not transcribe audio, so its content is unavailable]")
    except Exception as e:                                          # noqa: BLE001
        return None, f"[attachment '{name}' could not be read: {type(e).__name__}]"
    # mp4 / step / psd and anything else: 5 of 261 files across the benchmark
    return None, (f"[attachment '{name}' is a .{ext} file, which this suite cannot render; "
                  f"its content is unavailable]")


def _clip(text: str, name: str) -> str:
    text = (text or "").strip()
    if len(text) > _MAX_CHARS_PER_FILE:
        text = text[:_MAX_CHARS_PER_FILE] + f"\n[... '{name}' truncated at {_MAX_CHARS_PER_FILE} chars]"
    return text


def attachment_parts(rel_paths: List[str]) -> Tuple[List[Dict[str, Any]], int, int]:
    """Build the multimodal content parts for a task's attachments.

    Returns (parts, rendered, unavailable). Images go through the same lossless-base64
    path the vision suites use; everything else is rendered to text.
    """
    from .dataset_utils import extract_lossless_image_b64
    parts: List[Dict[str, Any]] = []
    rendered = unavailable = 0
    for rel in rel_paths or []:
        name = os.path.basename(rel)
        ext = rel.rsplit(".", 1)[-1].lower() if "." in rel else ""
        if ext in ("png", "jpg", "jpeg", "webp", "gif", "bmp"):
            local = _fetch(rel)
            b64 = None
            if local:
                try:
                    from PIL import Image
                    b64 = extract_lossless_image_b64(Image.open(local))
                except Exception as e:                              # noqa: BLE001
                    logger.warning("gdpval: image %s failed (%s)", name, e)
            if b64:
                parts.append({"type": "text", "text": f"\n--- attachment: {name} ---"})
                parts.append({"type": "image_url",
                              "image_url": {"url": f"data:image/png;base64,{b64}"}})
                rendered += 1
            else:
                parts.append({"type": "text",
                              "text": f"\n[attachment '{name}' could not be rendered]"})
                unavailable += 1
            continue
        text, note = render_attachment(rel)
        if text:
            parts.append({"type": "text", "text": f"\n--- attachment: {name} ---\n{text}"})
            rendered += 1
        else:
            parts.append({"type": "text", "text": "\n" + (note or f"[{name} unavailable]")})
            unavailable += 1
    return parts, rendered, unavailable


def render_reference(rel_paths: List[str]) -> Tuple[Optional[str], Optional[str]]:
    """Render the expert reference deliverable(s) to text for the pairwise comparison.

    Returns (text, note). The reference is a produced FILE (.xlsx/.docx/.pdf/...); it is
    rendered to text with the same converters the inputs use, so the model's prose answer is
    compared against the reference's rendered CONTENT. An image-only or unrenderable reference
    yields (None, note) and the task is excluded from the win-rate rather than scored as a loss.
    """
    if not rel_paths:
        return None, "[no reference deliverable shipped for this task]"
    chunks: List[str] = []
    notes: List[str] = []
    for rel in rel_paths:
        name = os.path.basename(rel)
        ext = rel.rsplit(".", 1)[-1].lower() if "." in rel else ""
        if ext in ("png", "jpg", "jpeg", "webp", "gif", "bmp"):
            notes.append(f"[reference '{name}' is an image; not rendered to text]")
            continue
        text, note = render_attachment(rel)
        if text:
            chunks.append(f"--- reference deliverable: {name} ---\n{text}")
        else:
            notes.append(note or f"[reference '{name}' unavailable]")
    if not chunks:
        return None, (" ".join(notes) if notes
                      else "[reference deliverable could not be rendered to text]")
    return "\n\n".join(chunks), (" ".join(notes) if notes else None)


# --------------------------------------------------------------------------- #
# loading
# --------------------------------------------------------------------------- #
def _load_gdpval_samples(
    limit: Optional[int] = None,
) -> Tuple[List[Tuple[List[Dict[str, Any]], str, Dict[str, Any]]], Dict[str, Dict[str, Any]]]:
    """Load GDPval tasks with inputs rendered and expert references resolved.

    Returns (samples, references) where references maps task_id -> {"text", "note", "n_files"}.
    The reference text is carried on a SIDE CHANNEL, not in the sample metadata, because
    base.run_eval_suite merges sample metadata into the request payload sent to the model - a
    20 KB reference deliverable does not belong in every generation request.
    """
    try:
        from datasets import load_dataset
        rows = list(load_dataset(_REPO, split="train"))
    except Exception as e:
        logger.error(f"Failed to load dataset for gdpval: {e}")
        raise RuntimeError(f"Could not load dataset for gdpval: {e}") from e
    if not rows:
        raise RuntimeError("Dataset for gdpval returned empty rows")

    # Stratified by occupation, not a contiguous head (audit RC-1).
    rows = stratified_sample(rows, limit, lambda r: (r or {}).get("occupation"), seed="gdpval")

    samples: List[Tuple[List[Dict[str, Any]], str, Dict[str, Any]]] = []
    references: Dict[str, Dict[str, Any]] = {}
    with_files = missing_any = with_reference = 0
    for item in rows:
        task_id = str(item.get("task_id") or "")
        prompt = str(item.get("prompt") or "").strip()
        if not task_id or not prompt:
            raise RuntimeError("gdpval: unexpected schema (task_id/prompt); "
                               "refusing to fabricate sample data")
        try:
            criteria = json.loads(item.get("rubric_json") or "[]")
        except Exception:
            criteria = []
        if not criteria:
            raise RuntimeError(f"gdpval: task {task_id} has no parseable rubric_json; "
                               "refusing to score against an empty rubric")

        refs = list(item.get("reference_files") or [])
        parts: List[Dict[str, Any]] = [{"type": "text", "text": prompt}]
        rendered = unavailable = 0
        if refs:
            with_files += 1
            extra, rendered, unavailable = attachment_parts(refs)
            parts += extra
            if unavailable:
                missing_any += 1
        content: Any = parts if len(parts) > 1 else prompt

        # Expert reference deliverable(s) for the pairwise comparison (side channel).
        deliverables = list(item.get("deliverable_files") or [])
        ref_text, ref_note = render_reference(deliverables)
        references[task_id] = {"text": ref_text, "note": ref_note, "n_files": len(deliverables)}
        if ref_text:
            with_reference += 1

        samples.append((
            [{"role": "user", "content": content}],
            json.dumps(criteria),
            {"category": str(item.get("occupation") or item.get("sector") or "knowledge_work"),
             "task_id": task_id,
             "attachments": len(refs),
             "attachments_rendered": rendered,
             "attachments_unavailable": unavailable},
        ))

    logger.info("Loaded %d gdpval samples (%d with input attachments; %d had an unrenderable "
                "input; %d have a text-renderable expert reference for the pairwise score).",
                len(samples), with_files, missing_any, with_reference)
    return samples, references


# --------------------------------------------------------------------------- #
# scoring: pairwise win-rate vs the expert reference (canonical GDPval metric)
# --------------------------------------------------------------------------- #
_PAIRWISE_PROMPT = """You are an expert grader for economically-valuable knowledge work.

You are given a task, its grading rubric, and TWO candidate deliverables, A and B. Their
order is arbitrary and carries no meaning - do not prefer a deliverable for being first.
Decide which deliverable better completes the task, judged against the SUBSTANCE of the
rubric: correctness, completeness, sound reasoning, and coverage of what the task asks for.
Both deliverables are shown as text (a file deliverable has been rendered to text), so judge
the CONTENT, not the file format, byte layout, or presentation.

# TASK
{prompt}

# GRADING RUBRIC
{rubric}

# DELIVERABLE A
{a}

# DELIVERABLE B
{b}

Reply with ONLY a JSON object:
{{"winner": "A", "reason": "<one short sentence>"}}
where "winner" is "A", "B", or "tie" (use "tie" only when the two are genuinely of equal quality)."""


def _task_prompt(trace: Dict[str, Any]) -> str:
    """The task text from the trace's first user message (string or multimodal parts)."""
    msgs = trace.get("messages") or []
    if not msgs:
        return ""
    content = msgs[0].get("content")
    if isinstance(content, str):
        return content
    if isinstance(content, list):
        for part in content:
            if isinstance(part, dict) and part.get("type") == "text":
                return str(part.get("text") or "")
    return ""


def _model_score(winner: str, model_position: str) -> float:
    """Score a single pairwise verdict from the MODEL's perspective.

    winner is the judge's "a"/"b"/"tie"; model_position is which slot the model occupied in
    that call ("a" or "b"). Win = 1.0, tie = 0.5, loss = 0.0. An unparseable/other winner is
    treated as a tie (0.5) rather than a win or a loss, so a malformed judge reply cannot
    swing the result in either direction.
    """
    w = (winner or "").strip().lower()
    if w == "tie":
        return 0.5
    if w in ("a", "b"):
        return 1.0 if w == model_position else 0.0
    return 0.5


def _make_scorer(judge_model: str, judge_concurrency: int, metrics: Dict[str, Any],
                 references: Dict[str, Dict[str, Any]]):
    """Grade each task by a position-swapped pairwise comparison vs the expert reference."""
    async def _score(sample_traces: List[Dict[str, Any]]) -> None:
        import asyncio
        from tqdm import tqdm

        sem = asyncio.Semaphore(max(1, judge_concurrency))

        async def _compare(prompt_text: str, rubric_text: str, a_text: str, b_text: str):
            """One pairwise judge call -> 'a' | 'b' | 'tie' | None (outage/unparseable-json)."""
            prompt = _PAIRWISE_PROMPT.format(
                prompt=prompt_text[:_MAX_JUDGE_CHARS],
                rubric=rubric_text[:_MAX_RUBRIC_CHARS],
                a=a_text[:_MAX_JUDGE_CHARS],
                b=b_text[:_MAX_JUDGE_CHARS])
            async with sem:
                text, _judge_used = await judge_generate_cascade(prompt)
            if text is None:
                return None
            m = re.search(r"\{[\s\S]*\}", text)
            if not m:
                return None
            try:
                obj = json.loads(m.group(0))
            except Exception:                                       # noqa: BLE001
                return None
            return str(obj.get("winner", "")).strip().lower()

        async def _one(tr: Dict[str, Any], pbar) -> None:
            ep = tr.get("extra_payload") or {}
            task_id = str(ep.get("task_id") or "")
            ref = references.get(task_id) or {}
            ref_text = ref.get("text")
            try:
                criteria = json.loads(tr.get("gold_answer") or "[]")
            except Exception:
                criteria = []
            rubric_text = "\n".join(
                f'[{c.get("score")} pts] {c.get("criterion")}' for c in criteria) or "(no rubric)"
            model_text = strip_thinking_tags(tr.get("response_text") or "").strip()
            tr["reference_excerpt"] = (ref_text or "")[:2000]
            tr["reference_note"] = ref.get("note")

            if not ref_text:
                # No text-renderable expert reference -> the pairwise metric is undefined for
                # this task. Exclude it (not a loss), so it cannot deflate the win-rate.
                tr["is_correct"], tr["status"] = False, "NO_REFERENCE"
                tr["scoring_excluded"] = True
                pbar.update(1)
                return
            if not model_text:
                # A failed/empty model deliverable loses the pairwise (it produced nothing),
                # rather than being dropped - a model cannot inflate its win-rate by failing.
                tr["pairwise_score"] = 0.0
                tr["pairwise_detail"] = {"empty_response": True}
                tr["is_correct"], tr["status"] = False, "OK"
                pbar.update(1)
                return

            prompt_text = _task_prompt(tr)
            # Position-swapped: model as A (ref as B), then model as B (ref as A). Averaging the
            # two cancels the judge's order bias.
            w1 = await _compare(prompt_text, rubric_text, model_text, ref_text)   # model = A
            w2 = await _compare(prompt_text, rubric_text, ref_text, model_text)   # model = B
            if w1 is None or w2 is None:
                # Judge cascade exhausted / unusable: excluded from the win-rate and from
                # base.run_eval_suite's pass/fail accuracy, NOT scored as a loss.
                tr["judge_grade"] = "JUDGE_OUTAGE"
                tr["status"] = "OK"
                pbar.update(1)
                return
            s1 = _model_score(w1, "a")
            s2 = _model_score(w2, "b")
            pairwise = (s1 + s2) / 2.0
            tr["pairwise_score"] = pairwise
            tr["pairwise_detail"] = {
                "call1_winner": w1, "call2_winner": w2, "consistent": s1 == s2}
            tr["is_correct"] = pairwise >= 0.5   # "won or tied" == not a loss
            tr["status"] = "OK"
            pbar.update(1)

        with tqdm(total=len(sample_traces), desc="Judging [GDPVAL pairwise]") as pbar:
            await asyncio.gather(*[_one(t, pbar) for t in sample_traces])

        eligible = [t for t in sample_traces
                    if t.get("status") == "OK" and t.get("judge_grade") != "JUDGE_OUTAGE"
                    and t.get("pairwise_score") is not None]
        scores = [t["pairwise_score"] for t in eligible]
        wins = sum(1 for s in scores if s == 1.0)
        losses = sum(1 for s in scores if s == 0.0)
        ties = len(scores) - wins - losses          # 0.5, or a 0.25/0.75 position split
        metrics["gdpval_report"] = {
            "win_rate": round(sum(scores) / len(scores), 4) if scores else None,
            "tasks_in_win_rate": len(scores),
            "wins": wins,
            "ties_or_split": ties,
            "losses": losses,
            "consistent_pairs": sum(
                1 for t in eligible if (t.get("pairwise_detail") or {}).get("consistent")),
            "tasks_no_reference": sum(
                1 for t in sample_traces if t.get("status") == "NO_REFERENCE"),
            "tasks_judge_outage": sum(
                1 for t in sample_traces if t.get("judge_grade") == "JUDGE_OUTAGE"),
            "tasks_with_unavailable_attachments": sum(
                1 for t in sample_traces
                if (t.get("extra_payload") or {}).get("attachments_unavailable")),
        }
    return _score


def run_gdpval(
    model_name: str,
    base_url: str,
    concurrency: int,
    enable_thinking: bool = False,
    **kwargs,
) -> Dict[str, Any]:
    """Execute gdpval: attachments in, pairwise win-rate vs the expert reference out.

    The headline `accuracy` is the pairwise WIN-RATE (wins + 0.5*ties) of the model's
    deliverable against the human-expert reference, position-swapped to cancel order bias -
    the canonical GDPval measure. It is NOT OpenAI's GDPval number (`leaderboard_comparable`
    False): a text endpoint cannot emit the .xlsx/.docx/.pdf deliverable, so the reference file
    is rendered to text and the model competes on rendered content only.
    """
    skip = gemini_required_skip("gdpval", model_name)
    if skip:
        return skip
    metrics: Dict[str, Any] = {}
    samples, references = _load_gdpval_samples(limit=kwargs.get("limit"))
    result = run_eval_suite(
        eval_name="gdpval",
        model_name=model_name,
        base_url=base_url,
        concurrency=concurrency,
        samples=samples,
        async_eval_fn=_make_scorer(
            kwargs.get("judge_model", DEFAULT_JUDGE_MODEL), concurrency, metrics, references),
        thinking=enable_thinking,
        extra_payload=kwargs.get("extra_payload"),
        limit=kwargs.get("limit"),
        max_output_tokens=kwargs.get("max_output_tokens"),
    )
    result.update(metrics)
    # A text endpoint cannot produce the FILE deliverable the reference is, so the model
    # competes on rendered content only - this is not OpenAI's GDPval number. Grading uses
    # gbench's standard Gemini cascade (a gbench convention), which is not the reason here.
    result["leaderboard_comparable"] = False
    result["leaderboard_comparable_reason"] = (
        "text endpoint vs an expert FILE deliverable rendered to text (the model cannot emit "
        ".xlsx/.docx/.pdf and competes on rendered content only), graded by gbench's standard "
        "Gemini cascade; a gbench-internal win-rate, not OpenAI's GDPval number")
    rep = metrics.get("gdpval_report") or {}

    win_rate = rep.get("win_rate")
    if win_rate is not None:
        result["accuracy"] = round(win_rate * 100.0, 2)   # headline = pairwise win-rate (%)
        result["win_rate_pct"] = result["accuracy"]
        scored = rep.get("tasks_in_win_rate") or 0
        # `total_questions` counts tasks ATTEMPTED; the win-rate is over tasks that could be
        # SCORED (a text-renderable reference + a live judge). Name the gap rather than hiding
        # it in a denominator: an excluded task was unmeasurable, not asked-and-lost.
        result["tasks_scored"] = scored
        no_ref = rep.get("tasks_no_reference") or 0
        outage = rep.get("tasks_judge_outage") or 0
        result["tasks_no_reference"] = no_ref
        result["tasks_judge_outage"] = outage
        total_q = result.get("total_questions") or 0
        result["tasks_unaccounted"] = max(0, total_q - scored - no_ref - outage)
        result["scoring_note"] = (
            f"accuracy is the PAIRWISE WIN-RATE (wins + 0.5*ties) of the model's deliverable vs "
            f"the expert reference over the {scored} task(s) with a text-renderable reference and "
            f"a live judge, position-swapped to cancel order bias; NOT a pass rate over "
            f"{total_q}. Tasks without a renderable reference ({no_ref}) or hit by a judge outage "
            f"({outage}) are excluded from the win-rate, not counted as losses; an empty model "
            f"deliverable IS counted, as a loss.")
        if no_ref or outage:
            logger.warning(
                "[gdpval] win-rate covers %d of %d task(s): %d had no text-renderable expert "
                "reference and %d hit a judge outage (both excluded, not scored as losses).",
                scored, total_q, no_ref, outage)
    logger.warning(
        "[gdpval] PARTIAL: pairwise win-rate vs an expert FILE deliverable rendered to text. A "
        "text endpoint cannot emit .xlsx/.docx/.pdf, so this is not comparable with OpenAI's "
        "published GDPval win-rate.")
    return result
